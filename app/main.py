"""Market Visit Tracker API with GTM hierarchy and retailer intelligence."""

import calendar
import csv
import io
import os
from datetime import date
from pathlib import Path
from typing import List, Optional

import httpx
from openpyxl import load_workbook

from fastapi import Depends, FastAPI, HTTPException, Query, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse, Response

from . import crud, database
from .auth import get_current_user, get_dashboard_user, hash_password, require_admin_key, require_manager, verify_password, create_token
from .schemas import (
    AdminUserOut,
    CoverageRow,
    LoginOut,
    PinLoginRequest,
    LoginRequest,
    PasswordResetOut,
    RetailerHealthOut,
    RetailerOut,
    StatsOut,
    UserOut,
    UserStatusOut,
    VisitCreate,
    VisitOut, PhotoOut,
)
app = FastAPI(title="Market Visit Tracker API", description="GTM retailer visits, coverage and field intelligence.", version="2.1.0")

VWORK_LIVE_API_URL = os.getenv("VWORK_LIVE_API_URL", "http://127.0.0.1:8000").rstrip("/")
VWORK_LIVE_API_KEY = os.getenv("VWORK_LIVE_API_KEY", "").strip()
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=False, allow_methods=["*"], allow_headers=["*"])

STATIC_DIR = Path(__file__).resolve().parent / "static"


@app.on_event("startup")
def on_startup():
    database.init_db()


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/", include_in_schema=False)
def frontend():
    return FileResponse(STATIC_DIR / "index.html")


@app.post("/auth/login", response_model=LoginOut)
def login(payload: LoginRequest):
    user = crud.get_user_by_username(payload.username)
    if not user or not user.get("password_hash") or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid username or password.")
    token = create_token(user)
    safe_user = {k: user.get(k) for k in ("id", "name", "role", "tl", "ss", "rds", "active")}
    return {"access_token": token, "token_type": "bearer", "user": safe_user}


@app.get("/auth/pin-users", include_in_schema=False)
def pin_users():
    rows = crud.list_users_admin()
    return {
        "TL": sorted([r["name"] for r in rows if r["role"] == "TL"]),
        "SS": sorted([r["name"] for r in rows if r["role"] == "SS"]),
    }


@app.post("/auth/pin-login", response_model=LoginOut)
def pin_login(payload: PinLoginRequest):
    user = crud.get_user_by_pin(payload.pin)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid PIN.")
    token = create_token(user)
    safe_user = {k: user.get(k) for k in ("id", "name", "role", "tl", "ss", "rds", "active")}
    return {"access_token": token, "token_type": "bearer", "user": safe_user}


@app.get("/auth/me", response_model=UserOut)
def me(user: dict = Depends(get_current_user)):
    if user.get("role") == "MANAGER":
        return {"id": 0, "name": "Manager", "role": "MANAGER", "active": True}
    db_user = crud.get_user_by_username(user.get("sub", user.get("name", "")))
    if not db_user:
        raise HTTPException(status_code=401, detail="User account not found.")
    return db_user


def _scope(user: dict) -> dict:
    role = user.get("role")
    name = user.get("name")

    if role == "TL":
        return {"tl": name}
    if role == "SS":
        return {"ss": name}
    if role == "RDS":
        return {"rds": name}

    return {}

@app.post("/visits", response_model=VisitOut)
def create_visit(payload: VisitCreate, user: dict = Depends(get_current_user)):
    scope = _scope(user)
    retailer = crud.lookup_retailer_by_name(payload.retailer)
    if not retailer:
        raise HTTPException(status_code=400, detail="Retailer is not in the master list.")
    if scope and any(retailer.get(k) != v for k, v in scope.items()):
        raise HTTPException(status_code=403, detail="You can only log visits for retailers assigned to you.")
    payload.submitted_by = user.get("name", "Manager")
    payload.submitted_role = user.get("role", "MANAGER")
    return crud.create_visit(payload)


@app.get("/visits", response_model=List[VisitOut])
def get_visits(retailer: Optional[str] = None, tl: Optional[str] = None, ss: Optional[str] = None,
               rds: Optional[str] = None, submitted_by: Optional[str] = None,
               submitted_role: Optional[str] = None, date_from: Optional[str] = Query(default=None),
               date_to: Optional[str] = Query(default=None), search: Optional[str] = Query(default=None),
               limit: int = Query(default=500, ge=1, le=2000), offset: int = Query(default=0, ge=0),
               user: dict = Depends(get_dashboard_user)):
    scope = _scope(user)
    if scope:
        tl, ss, rds = scope.get("tl"), scope.get("ss"), scope.get("rds")
        submitted_by = None
    return crud.list_visits(retailer=retailer, tl=tl, ss=ss, rds=rds, submitted_by=submitted_by,
                            submitted_role=submitted_role, date_from=date_from, date_to=date_to,
                            search=search, limit=limit, offset=offset)


@app.get("/retailers", response_model=List[RetailerOut])
def get_retailers(tl: Optional[str] = None, ss: Optional[str] = None, rds: Optional[str] = None,
                  search: Optional[str] = None, user: dict = Depends(get_dashboard_user)):
    scope = _scope(user)
    if scope:
        tl, ss, rds = scope.get("tl"), scope.get("ss"), scope.get("rds")
    return crud.list_retailers(tl=tl, ss=ss, rds=rds, search=search)


@app.get("/retailers/{retailer_code}/360")
def retailer_360(
    retailer_code: str,
    start_date: Optional[str] = Query(default=None, alias="startDate"),
    end_date: Optional[str] = Query(default=None, alias="endDate"),
    inventory_date: Optional[str] = Query(default=None, alias="inventoryDate"),
    user: dict = Depends(get_dashboard_user),
):
    retailer = crud.lookup_retailer_by_code(retailer_code)
    if not retailer:
        raise HTTPException(status_code=404, detail="Retailer is not in the Market Visit master.")

    scope = _scope(user)
    if scope and any(retailer.get(k) != v for k, v in scope.items()):
        raise HTTPException(status_code=403, detail="You cannot access this retailer.")

    today = date.today()
    start_date = start_date or today.replace(day=1).isoformat()
    end_date = end_date or today.isoformat()
    inventory_date = inventory_date or end_date

    params = {
        "startDate": start_date,
        "endDate": end_date,
        "inventoryDate": inventory_date,
        "pageSize": 500,
        "maxPages": 100,
    }
    headers = {}
    if VWORK_LIVE_API_KEY:
        headers["X-API-Key"] = VWORK_LIVE_API_KEY

    try:
        with httpx.Client(timeout=httpx.Timeout(240.0, connect=10.0)) as client:
            response = client.get(
                f"{VWORK_LIVE_API_URL}/api/v1/retailers/{retailer['code']}/360",
                params=params,
                headers=headers,
            )
    except httpx.HTTPError as exc:
        raise HTTPException(status_code=502, detail=f"Retailer 360 service unavailable: {exc}") from exc

    if response.status_code == 404:
        raise HTTPException(status_code=404, detail="No live V-Work data found for this retailer.")
    if not response.is_success:
        detail = "Live retailer intelligence request failed."
        try:
            detail = response.json().get("detail") or detail
        except Exception:
            pass
        raise HTTPException(status_code=502, detail=detail)

    live = response.json()
    visits = crud.retailer_visit_history(retailer["name"], limit=5)

    sales = live.get("sales") or {}
    inventory = live.get("inventory") or {}
    performance = live.get("performance") or {}
    sales_units = int(sales.get("units") or 0)
    stock_units = int(inventory.get("units") or 0)
    dos = performance.get("dos")

    target_month = start_date[:7]
    target_row = crud.get_retailer_target(retailer["code"], target_month)
    target_volume = int(target_row.get("target_volume") or 0) if target_row else None
    target_value = int(target_row.get("target_value") or 0) if target_row else None
    achievement_pct = None
    gap_volume = None
    required_run_rate = None
    days_remaining = None
    if target_volume is not None:
        achievement_pct = round((sales_units / target_volume) * 100, 1) if target_volume > 0 else 0
        gap_volume = max(target_volume - sales_units, 0)
        period_end = date.fromisoformat(end_date)
        month_days = calendar.monthrange(period_end.year, period_end.month)[1]
        days_remaining = max(month_days - period_end.day + 1, 1)
        required_run_rate = round(gap_volume / days_remaining, 1)

    actions = []
    if sales_units == 0:
        actions.append("Discuss sell-out activation and identify why MTD sales are zero.")
    if isinstance(dos, (int, float)):
        if dos < 7:
            actions.append("Stock cover is below 7 days. Prioritise replenishment on fast-moving models.")
        elif dos > 30:
            actions.append("Stock cover is above 30 days. Focus on ageing stock and sell-through actions.")
    model_intelligence = performance.get("model_intelligence") or []
    critical_models = [
        x for x in model_intelligence
        if x.get("status") in ("Out of Stock", "Low Stock", "High Stock", "No MTD Sell-out")
    ]

    for item in critical_models[:4]:
        model = item.get("model") or "Unknown model"
        status = item.get("status")
        sales_qty = int(item.get("sales") or 0)
        stock_qty = int(item.get("stock") or 0)
        model_dos = item.get("dos")
        if status == "Out of Stock":
            actions.append(f"{model}: {sales_qty} MTD sales but zero stock. Replenish immediately.")
        elif status == "Low Stock":
            actions.append(f"{model}: only {model_dos} DOS with {sales_qty} MTD sales. Prioritise replenishment.")
        elif status == "High Stock":
            actions.append(f"{model}: high stock cover at {model_dos} DOS. Discuss sell-through activity.")
        elif status == "No MTD Sell-out":
            actions.append(f"{model}: {stock_qty} units in stock with no MTD sell-out. Identify conversion blocker.")

    if not visits:
        actions.append("No prior market visit is recorded. Capture retailer feedback, commitments and follow-up actions.")
    else:
        last_visit = visits[0]
        if last_visit.get("suggestions"):
            actions.append("Review previous follow-up: " + str(last_visit["suggestions"])[:180])

    return {
        **live,
        "target": {
            "month": target_month,
            "volume": target_volume,
            "value": target_value,
            "achievement_pct": achievement_pct,
            "gap_volume": gap_volume,
            "required_run_rate": required_run_rate,
            "days_remaining": days_remaining,
            "loaded": target_row is not None,
        },
        "market_visit": {
            "master": retailer,
            "recent_visits": visits,
            "what_to_discuss": actions[:5],
        },
    }


@app.get("/visits/{visit_id}", response_model=VisitOut)
def get_visit(visit_id: int, user: dict = Depends(get_dashboard_user)):
    visit = crud.get_visit(visit_id)
    if not visit:
        raise HTTPException(status_code=404, detail="Visit not found.")
    scope = _scope(user)
    if scope and any(visit.get(k) != v for k, v in scope.items()):
        raise HTTPException(status_code=403, detail="You cannot access this visit.")
    return visit


@app.get("/visits/{visit_id}/photos", response_model=List[PhotoOut])
def visit_photos(visit_id: int, user: dict = Depends(get_dashboard_user)):
    visit = crud.get_visit(visit_id)
    if not visit:
        raise HTTPException(status_code=404, detail="Visit not found.")
    scope = _scope(user)
    if scope and any(visit.get(k) != v for k, v in scope.items()):
        raise HTTPException(status_code=403, detail="You cannot access this visit.")
    rows = crud.list_visit_photos(visit_id)
    return [{**r, "url": f"/visits/{visit_id}/photos/{r['id']}"} for r in rows]

@app.get("/visits/{visit_id}/photos/{photo_id}")
def visit_photo(visit_id: int, photo_id: int, user: dict = Depends(get_dashboard_user)):
    visit = crud.get_visit(visit_id)
    photo = crud.get_visit_photo(photo_id)
    if not visit or not photo or photo["visit_id"] != visit_id:
        raise HTTPException(status_code=404, detail="Photo not found.")
    scope = _scope(user)
    if scope and any(visit.get(k) != v for k, v in scope.items()):
        raise HTTPException(status_code=403, detail="You cannot access this photo.")
    return Response(content=photo["data"], media_type=photo["mime_type"])

@app.post("/visits/{visit_id}/photos", response_model=List[PhotoOut])
async def upload_visit_photos(visit_id: int, files: List[UploadFile] = File(...), user: dict = Depends(get_current_user)):
    visit = crud.get_visit(visit_id)
    if not visit:
        raise HTTPException(status_code=404, detail="Visit not found.")
    scope = _scope(user)
    if scope and any(visit.get(k) != v for k, v in scope.items()):
        raise HTTPException(status_code=403, detail="You cannot upload photos to this visit.")
    existing = crud.list_visit_photos(visit_id)
    if len(existing) + len(files) > 5:
        raise HTTPException(status_code=400, detail="Maximum 5 stock photos per visit.")
    allowed = {"image/jpeg", "image/png", "image/webp"}
    out = []
    for upload in files:
        if upload.content_type not in allowed:
            raise HTTPException(status_code=400, detail="Only JPG, PNG or WEBP images are allowed.")
        data = await upload.read()
        if len(data) > 4 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="Each photo must be 4 MB or smaller.")
        row = crud.add_visit_photo(visit_id, upload.filename or "stock-photo", upload.content_type, data)
        out.append({**row, "url": f"/visits/{visit_id}/photos/{row['id']}"})
    return out

@app.delete("/visits/{visit_id}", status_code=204, dependencies=[Depends(require_admin_key)])
def delete_visit(visit_id: int):
    if not crud.delete_visit(visit_id):
        raise HTTPException(status_code=404, detail="Visit not found.")
    return None


@app.get("/visits/export.csv")
def export_visits_csv(tl: Optional[str] = None, ss: Optional[str] = None, rds: Optional[str] = None,
                      date_from: Optional[str] = None, date_to: Optional[str] = None,
                      user: dict = Depends(get_dashboard_user)):
    scope = _scope(user)
    if scope: tl, ss, rds = scope.get("tl"), scope.get("ss"), scope.get("rds")
    rows = crud.list_visits(tl=tl, ss=ss, rds=rds, date_from=date_from, date_to=date_to, limit=2000)
    buffer = io.StringIO(); writer = csv.writer(buffer)
    writer.writerow(["id", "visit_date", "retailer", "tl", "ss", "rds", "market", "feedback", "suggestions", "submitted_by", "submitted_role", "created_at"])
    for r in rows:
        writer.writerow([r["id"], r["visit_date"], r["retailer"], r["tl"], r["ss"], r["rds"], r["market"], r["feedback"], r["suggestions"], r["submitted_by"], r["submitted_role"], r["created_at"]])
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="text/csv", headers={"Content-Disposition": "attachment; filename=market-visits.csv"})


@app.get("/users", response_model=List[UserOut])
def get_users(role: Optional[str] = Query(default=None, pattern="^(TL|SS|RDS)$"), user: dict = Depends(get_current_user)):
    require_manager(user)
    return crud.list_users(role=role)


@app.post("/admin/targets/upload", dependencies=[Depends(require_admin_key)])
async def admin_upload_targets(
    month: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    file: UploadFile = File(...),
):
    """Import monthly retailer targets from the standard Zone A target workbook."""
    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload an .xlsx target workbook.")

    try:
        raw = await file.read()
        workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        sheet = workbook[workbook.sheetnames[0]]
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to read target workbook: {exc}") from exc

    required = {"Retailer Code", "Retailer Name", "Vol Target", "Val Target"}
    header_row = None
    headers = None
    for row_no, values in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        candidate = [str(v).strip() if v is not None else "" for v in values]
        if required.issubset(set(candidate)):
            header_row = row_no
            headers = candidate
            break

    if header_row is None or headers is None:
        raise HTTPException(
            status_code=400,
            detail="Target workbook must contain Retailer Code, Retailer Name, Vol Target and Val Target columns.",
        )

    idx = {name: headers.index(name) for name in required}
    aggregated = {}
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        code = values[idx["Retailer Code"]] if idx["Retailer Code"] < len(values) else None
        if code is None or str(code).strip() == "":
            continue
        code = str(code).strip().upper()
        name = values[idx["Retailer Name"]] if idx["Retailer Name"] < len(values) else ""
        volume = values[idx["Vol Target"]] if idx["Vol Target"] < len(values) else 0
        value = values[idx["Val Target"]] if idx["Val Target"] < len(values) else 0
        try:
            volume = int(round(float(volume or 0)))
        except (TypeError, ValueError):
            volume = 0
        try:
            value = int(round(float(value or 0)))
        except (TypeError, ValueError):
            value = 0

        item = aggregated.setdefault(code, {
            "retailer_code": code,
            "retailer_name": str(name or "").strip(),
            "target_volume": 0,
            "target_value": 0,
        })
        item["target_volume"] += volume
        item["target_value"] += value
        if not item["retailer_name"] and name:
            item["retailer_name"] = str(name).strip()

    if not aggregated:
        raise HTTPException(status_code=400, detail="No retailer target rows were found.")

    result = crud.upsert_retailer_targets(month, list(aggregated.values()))
    return {
        **result,
        "filename": file.filename,
        "message": "Monthly retailer targets imported successfully.",
    }


@app.get("/admin/users",
    response_model=List[AdminUserOut],
    dependencies=[Depends(require_admin_key)],
)
def admin_users():
    return crud.list_users_admin()


@app.post(
    "/admin/users/{user_id}/reset-password",
    response_model=PasswordResetOut,
    dependencies=[Depends(require_admin_key)],
)
def admin_reset_password(user_id: int):
    result = crud.reset_user_password(user_id)

    if not result:
        raise HTTPException(
            status_code=404,
            detail="User not found.",
        )

    return result


@app.post(
    "/admin/users/{user_id}/activate",
    response_model=UserStatusOut,
    dependencies=[Depends(require_admin_key)],
)
def admin_activate_user(user_id: int):
    result = crud.set_user_status(
        user_id,
        True,
    )

    if not result:
        raise HTTPException(
            status_code=404,
            detail="User not found.",
        )

    return result


@app.post(
    "/admin/users/{user_id}/deactivate",
    response_model=UserStatusOut,
    dependencies=[Depends(require_admin_key)],
)
def admin_deactivate_user(user_id: int):
    result = crud.set_user_status(
        user_id,
        False,
    )

    if not result:
        raise HTTPException(
            status_code=404,
            detail="User not found.",
        )

    return result

@app.get("/stats", response_model=StatsOut)
def get_stats(user: dict = Depends(get_dashboard_user)):
    scope = _scope(user)
    return crud.get_stats(**scope)


@app.get("/coverage", response_model=List[CoverageRow])
def get_coverage(
    group_by: str = Query("tl", pattern="^(tl|ss|rds)$"),
    user: dict = Depends(get_dashboard_user)
):
    role = user.get("role")

    if role in ("TL", "SS", "RDS"):
        group_by = role.lower()

    rows = crud.get_coverage(group_by)

    if role in ("TL", "SS", "RDS"):
        name = user.get("name")
        rows = [r for r in rows if r["name"] == name]

    return rows


@app.get("/my-target-performance")
def my_target_performance(user: dict = Depends(get_dashboard_user)):
    """TL/SS monthly target vs achievement using assigned retailer targets and live V-Work MTD sales."""
    role = user.get("role")
    if role not in ("TL", "SS"):
        return {
            "available": False,
            "role": role,
            "name": user.get("name"),
            "message": "Target performance is available for TL and SS users.",
        }

    today = date.today()
    month = today.strftime("%Y-%m")
    scope = _scope(user)
    target = crud.get_hierarchy_target_summary(
        month,
        tl=scope.get("tl"),
        ss=scope.get("ss"),
    )
    retailer_codes = set(target.get("retailer_codes") or [])

    params = {
        "startDate": today.replace(day=1).isoformat(),
        "endDate": today.isoformat(),
        "inventoryDate": today.isoformat(),
        "pageSize": 500,
        "maxPages": 100,
    }
    headers = {}
    if VWORK_LIVE_API_KEY:
        headers["X-API-Key"] = VWORK_LIVE_API_KEY

    try:
        with httpx.Client(timeout=httpx.Timeout(240.0, connect=10.0)) as client:
            response = client.get(
                f"{VWORK_LIVE_API_URL}/api/v1/retailers/priority-data",
                params=params,
                headers=headers,
            )
        response.raise_for_status()
        live_rows = (response.json() or {}).get("retailers") or []
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Live target achievement unavailable: {exc}") from exc

    achievement_units = sum(
        int(row.get("mtd_sales") or 0)
        for row in live_rows
        if str(row.get("retailer_code") or "").strip().upper() in retailer_codes
    )
    target_volume = int(target.get("target_volume") or 0)
    achievement_pct = round((achievement_units / target_volume) * 100, 1) if target_volume > 0 else 0
    gap = max(target_volume - achievement_units, 0)
    month_days = calendar.monthrange(today.year, today.month)[1]
    days_remaining = max(month_days - today.day + 1, 1)
    required_per_day = round(gap / days_remaining, 1) if target_volume > 0 else 0

    return {
        "available": True,
        "month": month,
        "role": role,
        "name": user.get("name"),
        "target": target_volume,
        "achievement": achievement_units,
        "achievement_pct": achievement_pct,
        "gap": gap,
        "required_per_day": required_per_day,
        "days_remaining": days_remaining,
        "retailers": len(retailer_codes),
    }


@app.get("/visit-priorities")
def visit_priorities(
    limit: int = Query(default=10, ge=1, le=25),
    user: dict = Depends(get_dashboard_user),
):
    """Rank retailers for today's field visit using recency + live sales velocity + Good Phone DOS."""
    scope = _scope(user)
    health_rows = crud.get_retailer_health(
        tl=scope.get("tl"),
        ss=scope.get("ss"),
        rds=scope.get("rds"),
        limit=2000,
    )

    today = date.today()
    params = {
        "startDate": today.replace(day=1).isoformat(),
        "endDate": today.isoformat(),
        "inventoryDate": today.isoformat(),
        "pageSize": 500,
        "maxPages": 100,
    }
    headers = {}
    if VWORK_LIVE_API_KEY:
        headers["X-API-Key"] = VWORK_LIVE_API_KEY

    try:
        with httpx.Client(timeout=httpx.Timeout(240.0, connect=10.0)) as client:
            response = client.get(
                f"{VWORK_LIVE_API_URL}/api/v1/retailers/priority-data",
                params=params,
                headers=headers,
            )
        response.raise_for_status()
        live_rows = (response.json() or {}).get("retailers") or []
    except Exception:
        live_rows = []

    live_by_code = {
        str(r.get("retailer_code") or "").strip().upper(): r
        for r in live_rows
        if r.get("retailer_code")
    }

    ranked = []
    for row in health_rows:
        code = str(row.get("code") or "").strip().upper()
        live = live_by_code.get(code, {})
        days = row.get("days_since_visit")
        avg_daily = float(live.get("avg_daily_sales") or 0)
        dos = live.get("dos")
        stock = int(live.get("good_phone_stock") or 0)
        sales = int(live.get("mtd_sales") or 0)

        if days is None:
            recency_score = 50
        elif days >= 30:
            recency_score = 50
        elif days >= 15:
            recency_score = 40
        elif days >= 8:
            recency_score = 25
        else:
            recency_score = min(20, max(0, int(days) * 2))

        if sales > 0 and stock == 0:
            dos_score = 30
            stock_signal = "Out of stock"
        elif isinstance(dos, (int, float)) and dos < 7:
            dos_score = 30
            stock_signal = "Low stock"
        elif isinstance(dos, (int, float)) and dos > 45:
            dos_score = 25
            stock_signal = "High stock"
        elif isinstance(dos, (int, float)) and dos > 30:
            dos_score = 15
            stock_signal = "Elevated stock"
        else:
            dos_score = 0
            stock_signal = "Stock balanced"

        if avg_daily >= 5:
            velocity_score = 20
        elif avg_daily >= 2:
            velocity_score = 15
        elif avg_daily >= 1:
            velocity_score = 10
        elif avg_daily > 0:
            velocity_score = 5
        else:
            velocity_score = 0

        reasons = []
        if days is None:
            reasons.append("Never visited")
        elif days >= 15:
            reasons.append(f"{days} days since last visit")
        elif days >= 8:
            reasons.append(f"Follow-up due after {days} days")
        if stock_signal != "Stock balanced":
            reasons.append(stock_signal)
        if avg_daily >= 2:
            reasons.append(f"Strong velocity {avg_daily:.1f}/day")
        elif sales == 0:
            reasons.append("No MTD sales")

        ranked.append({
            "retailer_code": code,
            "retailer": row.get("name"),
            "zone": row.get("zone"),
            "club": row.get("club"),
            "rds": row.get("rds"),
            "last_visit": row.get("last_visit"),
            "days_since_visit": days,
            "mtd_sales": sales,
            "avg_daily_sales": avg_daily,
            "good_phone_stock": stock,
            "dos": dos,
            "priority_score": recency_score + dos_score + velocity_score,
            "reason": " · ".join(reasons[:3]) or "Routine coverage",
        })

    ranked.sort(key=lambda x: (x["priority_score"], x["avg_daily_sales"]), reverse=True)
    return ranked[:limit]


@app.get("/retailer-health", response_model=List[RetailerHealthOut])
def retailer_health(tl: Optional[str] = None, ss: Optional[str] = None, rds: Optional[str] = None,
                    zone: Optional[str] = None, priority: Optional[str] = None,
                    limit: int = Query(default=500, ge=1, le=2000), user: dict = Depends(get_dashboard_user)):
    scope = _scope(user)
    if scope: tl, ss, rds = scope.get("tl"), scope.get("ss"), scope.get("rds")
    return crud.get_retailer_health(tl=tl, ss=ss, rds=rds, zone=zone, priority=priority, limit=limit)
