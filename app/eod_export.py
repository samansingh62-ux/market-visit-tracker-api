"""Management-outcome Excel report for saved TL and SS EOD remarks."""

import io
import re
from collections import Counter, defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY="17365D"; BLUE="1F4E78"; LIGHT_BLUE="D9EAF7"; HEADER_BLUE="5B9BD5"
GREEN="70AD47"; PALE_GREEN="E2F0D9"; PALE_RED="F4CCCC"; RED="C00000"
PALE_ORANGE="FCE4D6"; PALE_YELLOW="FFF2CC"; WHITE="FFFFFF"; DARK="1F1F1F"

def _safe(v):
    if v is None: return ""
    if isinstance(v,(int,float)): return v
    s=str(v)
    return "'"+s if s.startswith(("=","+","-","@")) else s

def _norm(s): return re.sub(r"\s+"," ",str(s or "").strip()).lower()
def _any(t, terms): return any(x in t for x in terms)

def _issue(remark):
    t=_norm(remark)
    if _any(t,("outlet closed","retail closed","wod closed","close this wod","stopped billing","pivoting his business","owner is into other business","hardly open")):
        return "Structural / Closure"
    comp=_any(t,("realme","oppo","samsung","motorola","iphone","poco","competition","competitor","thailand trip","other company"))
    manpower=_any(t,("vba","fvba","manpower","promoter","shop boy","staff","resign","salary","id creat","id pending","work off","w/o"))
    if comp and manpower: return "Competition / Manpower"
    if manpower: return "Manpower / VBA"
    if _any(t,("scheme","offer","price","trip","focus on realme","focus is on realme")): return "Competition / Scheme"
    if _any(t,("payment issue","credit issue","payment","credit")): return "Commercial / Payment"
    if _any(t,("road construction","expansion of road","renovation","one way road","road expansion","construction")): return "Store / Access Disruption"
    if _any(t,("footfall","customer walking","less customer","few customer","market is down","market down","remote area","village area")): return "Market / Footfall"
    if _any(t,("overall sale is low","overall sale is not good","sale is not good","sales is low","counter sale is low","sale will come","sales will come","sale will happen","overall sales is low")): return "Generic Low Sale"
    return "Other"

def _has_action(r):
    t=_norm(r)
    return _any(t,("will ","action plan","asked","ask ","follow up","following up","discuss","talk","place ","apply ","social media","promotion","billing","sale return","close ","commitment","replacement","increase","focus"))

def _quantified(r):
    return bool(re.search(r"\b\d+\b",str(r or ""))) or _any(_norm(r),("pcs","pc ","pics","lakh","month"))

def _recent(last_visit, remark_date):
    if not last_visit: return False
    try:
        return 0 <= (date.fromisoformat(str(remark_date)[:10])-date.fromisoformat(str(last_visit)[:10])).days <= 1
    except Exception:
        return False

def _evaluate(row):
    remark=str(row.get("remark") or "").strip()
    issue=_issue(remark); target=int(row.get("target_volume") or 0)
    recent=_recent(row.get("last_visit"),row.get("remark_date")); action=_has_action(remark); quantified=_quantified(remark)
    if issue=="Structural / Closure" and target==0:
        return dict(issue=issue,evidence="Aligned: zero-target/closed or inactive outlet",score=90,
                    verdict="Strongly justified — structural/data hygiene",action=True,recent=recent,
                    justification="Inactive/closed outlet with zero target; this is a master-data/WOD hygiene issue rather than a normal daily sales issue.",
                    recommendation="Master-data action: close/reclassify WOD and remove/confirm target. Do not keep requesting recurring low-sale remarks.")
    score=(30 if issue not in ("Generic Low Sale","Other") else 10 if issue=="Other" else 0)
    score += 15 if quantified else 0; score += 20 if action else 0; score += 15 if recent else 0; score += 5 if target>0 else 0
    if issue=="Generic Low Sale": score-=10
    if _any(_norm(remark),("sale will come","sales will come","will happen in few days")): score-=5
    if len(remark)<28: score-=7
    score=max(0,min(90,score))
    verdict="Strongly justified" if score>=70 else "Partially justified" if score>=45 else "Weak / basic explanation"
    evidence="Specific root cause with usable evidence/action" if score>=70 else "Plausible, but evidence/action is incomplete" if score>=45 else "Not sufficiently justified"
    just={
      "Generic Low Sale":"Remark mostly restates low sales rather than identifying a root cause. It needs retailer-specific evidence explaining why sales are low.",
      "Manpower / VBA":"The remark identifies a manpower constraint affecting conversion/coverage. It should be supported by verified absence/resignation status and recovery timeline.",
      "Competition / Manpower":"The remark identifies competitor pressure plus manpower influence. It needs a quantified counter-plan.",
      "Competition / Scheme":"The remark identifies a competitor/scheme driver rather than a generic excuse. Management needs the commercial gap and retailer commitment.",
      "Market / Footfall":"Low footfall can be genuine but is frequently used as a generic explanation. Support it with visit evidence or observed outlet/competition sales.",
      "Store / Access Disruption":"Road/renovation/access disruption is credible if duration and interim selling plan are documented.",
      "Commercial / Payment":"A payment/credit constraint can block billing. Include blocker details and resolution date.",
      "Structural / Closure":"The outlet may be structurally inactive or unsuitable and needs master-data validation."
    }.get(issue,"The remark has some retailer-specific context but needs stronger evidence and a measurable next step.")
    rec={
      "Generic Low Sale":"Return for revision: low sales is the symptom, not the cause. Require root cause, supporting evidence and measurable next action.",
      "Manpower / VBA":"Validate manpower status and replacement/attendance ETA; link the gap to expected sales recovery.",
      "Competition / Manpower":"Validate promoter deployment and competitor impact; set a retailer-specific counter-plan.",
      "Competition / Scheme":"Capture competitor offer/price benefit, share impact and the vivo counter-plan.",
      "Market / Footfall":"Validate footfall with visit evidence, outlet total sales or market observation; define activation plan.",
      "Store / Access Disruption":"Validate disruption duration and define an interim selling/activation plan.",
      "Commercial / Payment":"Confirm payment/credit blocker, owner and clearance ETA; track billing after resolution.",
      "Structural / Closure":"Validate outlet/WOD status and either reactivate with a plan or close/reclassify it."
    }.get(issue,"Accept provisionally and track the promised action to closure.")
    return dict(issue=issue,evidence=evidence,score=score,verdict=verdict,action=action,recent=recent,justification=just,recommendation=rec)

def _outcome(item):
    issue=item["issue"]; verdict=item["verdict"]; target=item["target_volume"]; score=item["score"]
    if issue=="Structural / Closure": decision="Close / Reclassify Outlet"
    elif verdict.startswith("Weak"): decision="Return for Rework"
    elif issue in ("Manpower / VBA","Competition / Manpower"): decision="Validate & Fix Manpower"
    elif issue=="Competition / Scheme": decision="Counter Competition / Scheme"
    elif issue=="Store / Access Disruption": decision="Validate Temporary Disruption"
    elif verdict.startswith("Strongly justified"): decision="Accept & Track"
    else: decision="Accept with Validation"
    if decision=="Close / Reclassify Outlet": priority="High" if target>0 else "Medium"
    elif decision=="Return for Rework": priority="Critical" if target>=50 or score<20 else "High"
    elif decision=="Validate & Fix Manpower" and target>=20: priority="High"
    elif target>=50: priority="High"
    elif target>=20: priority="Medium"
    else: priority="Low"
    if decision=="Close / Reclassify Outlet":
        owner="Channel Ops / Master Data"; next_action="Validate outlet status; close/reclassify WOD and remove/confirm target in master data."
    elif decision=="Validate & Fix Manpower":
        owner="TL / SS + Retail HR"; next_action="Confirm VBA/FVBA status, replacement/attendance ETA, and quantify expected sales recovery."
    elif decision=="Counter Competition / Scheme":
        owner="KAM / Channel + Scheme Team"; next_action="Capture competitor offer/share impact and submit a counter-plan with retailer commitment."
    elif decision=="Validate Temporary Disruption":
        owner="TL / SS"; next_action="Validate disruption duration and define interim selling/activation plan."
    elif decision=="Return for Rework":
        owner=item["submitted_by"] or item["tl"] or "TL / SS"; next_action="Re-submit with retailer-specific root cause, supporting evidence, and measurable next action."
    else:
        owner=item["submitted_by"] or "TL / SS"; next_action=item["recommendation"]
    days=0 if priority=="Critical" else 1 if priority=="High" else 3 if priority=="Medium" else 5
    due=date.today()+timedelta(days=days)
    return decision,priority,next_action,owner,due

def _style_header(ws,row=1):
    for c in ws[row]:
        c.fill=PatternFill("solid",fgColor=BLUE); c.font=Font(bold=True,color=WHITE)
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)

def _set_widths(ws,widths):
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

def _add_table(ws,name):
    if ws.max_row>=2:
        tab=Table(displayName=name,ref=ws.dimensions)
        tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showFirstColumn=False,showLastColumn=False,showColumnStripes=False)
        ws.add_table(tab)

def _build_items(rows):
    items=[]
    for r in rows:
        ev=_evaluate(r)
        item={**r,**ev}
        item.update(target_volume=int(r.get("target_volume") or 0),target_value=int(r.get("target_value") or 0))
        decision,priority,next_action,owner,due=_outcome(item)
        item.update(decision=decision,priority=priority,next_action=next_action,owner=owner,due=due)
        items.append(item)
    return items

def build_eod_workbook(rows,start,end,coverage=None):
    items=_build_items(rows)
    wb=Workbook()
    ws=wb.active; ws.title="Management Outcome"
    ws.merge_cells("A1:K2"); ws["A1"]=f"EOD Remarks — Management Outcome | {start.isoformat()}" if start==end else f"EOD Remarks — Management Outcome | {start.isoformat()} to {end.isoformat()}"
    ws["A1"].fill=PatternFill("solid",fgColor=NAVY); ws["A1"].font=Font(bold=True,color=WHITE,size=18); ws["A1"].alignment=Alignment(vertical="center")
    ws.merge_cells("A3:K3"); ws["A3"]="Outcome view: genuine root causes vs generic explanations, converted into management decisions and actions."
    ws["A3"].fill=PatternFill("solid",fgColor=LIGHT_BLUE); ws["A3"].font=Font(italic=True); ws["A3"].alignment=Alignment(wrap_text=True)
    total=len(items); strong=sum(i["verdict"].startswith("Strongly justified") for i in items); partial=sum(i["verdict"]=="Partially justified" for i in items); weak=sum(i["verdict"].startswith("Weak") for i in items); ready=sum(i["action"] for i in items); intervention=sum(i["priority"] in ("Critical","High") for i in items)
    kpis=[("A5","Total Remarks",total),("C5","Strongly Justified",strong),("E5","Partially Justified",partial),("G5","Weak / Basic",weak),("I5","Action-ready Remarks",ready),("K5","Needs Intervention",intervention)]
    for cell,label,val in kpis:
        ws[cell]=label; ws[cell].fill=PatternFill("solid",fgColor="D9E1F2" if label!="Needs Intervention" else PALE_RED); ws[cell].font=Font(bold=True,color=DARK if label!="Needs Intervention" else RED)
        vcell=f"{cell[0]}6"; ws[vcell]=val; ws[vcell].font=Font(bold=True,size=16,color=NAVY if label!="Needs Intervention" else RED)
    ws.merge_cells("A8:K8"); ws["A8"]="Management Conclusion"; ws["A8"].fill=PatternFill("solid",fgColor=HEADER_BLUE); ws["A8"].font=Font(bold=True,color=WHITE,size=13)
    ws.merge_cells("A9:K12")
    ws["A9"]="The EOD remarks are not uniformly decision-grade. Management should reject generic low-sale/footfall statements without evidence, validate manpower and competition claims, clean closed WODs from the target universe, and track accepted issues to a named owner and due date."
    ws["A9"].alignment=Alignment(wrap_text=True,vertical="top")
    ws.append([])
    ws["A14"]="Management Outcome"; ws["B14"]="Count"; ws["C14"]="% of Remarks"; ws["E14"]="Issue Category"; ws["F14"]="Count"; ws["G14"]="Management Interpretation"
    _style_header(ws,14)
    decisions=["Return for Rework","Validate & Fix Manpower","Accept with Validation","Accept & Track","Close / Reclassify Outlet","Counter Competition / Scheme","Validate Temporary Disruption"]
    dc=Counter(i["decision"] for i in items)
    for n,d in enumerate(decisions,15):
        ws[f"A{n}"]=d; ws[f"B{n}"]=dc[d]; ws[f"C{n}"]=dc[d]/total if total else 0; ws[f"C{n}"].number_format="0.0%"
    interp={
      "Generic Low Sale":"Symptom only; require root cause + evidence.",
      "Manpower / VBA":"Actionable if manpower status and ETA are verified.",
      "Competition / Manpower":"Validate promoter deployment and competitor impact.",
      "Competition / Scheme":"Requires counter-offer, retailer commitment and share impact.",
      "Structural / Closure":"Master-data/WOD hygiene issue, not a daily sales excuse.",
      "Store / Access Disruption":"Credible if duration and alternate action are documented.",
      "Market / Footfall":"Needs evidence; often overused as a generic explanation.",
      "Commercial / Payment":"Confirm blocker and clearance ETA.",
      "Other":"Needs case-specific validation and measurable next step."
    }
    ic=Counter(i["issue"] for i in items)
    for n,issue in enumerate(sorted(ic),15):
        ws[f"E{n}"]=issue; ws[f"F{n}"]=ic[issue]; ws[f"G{n}"]=interp.get(issue,"Needs management validation."); ws[f"G{n}"].alignment=Alignment(wrap_text=True)
    startrow=max(26,15+len(ic)+2)
    ws.merge_cells(start_row=startrow,start_column=1,end_row=startrow,end_column=11); ws.cell(startrow,1,"Priority Management Actions"); ws.cell(startrow,1).fill=PatternFill("solid",fgColor=GREEN); ws.cell(startrow,1).font=Font(bold=True,color=WHITE,size=13)
    actions=[
      ("1","Return weak remarks for rework","Low sales/footfall alone is not a root cause. Require evidence + specific action + ETA."),
      ("2","Close manpower gaps","Track VBA/FVBA resignation/absence cases with replacement owner and ETA."),
      ("3","Clean closed / inactive WODs","Remove/reclassify closed outlets and avoid recurring EOD explanations against invalid targets."),
      ("4","Challenge footfall claims","Require visit evidence, observed outlet sales/competition or market disruption proof."),
      ("5","Track competition / scheme cases","Capture competitor scheme, price/benefit gap, retailer focus and counter-plan."),
      ("6","Use the Action Tracker daily","Management Outcome + Priority + Owner + Due Date should drive next-day follow-up.")
    ]
    for off,(num,title,desc) in enumerate(actions,1):
        r=startrow+off; ws.cell(r,1,num); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); ws.cell(r,2,title).font=Font(bold=True); ws.merge_cells(start_row=r,start_column=5,end_row=r,end_column=11); ws.cell(r,5,desc).alignment=Alignment(wrap_text=True)
    _set_widths(ws,[24,12,14,4,24,10,48,4,18,12,18]); ws.freeze_panes="A4"

    # Action Tracker
    at=wb.create_sheet("Action Tracker")
    headers=["ID","Date","Retailer Code","Retailer","Role","Submitted By","KAM","SS","TL","RDS","Club","Town","District","Target Vol","EOD Remark","Issue Category","Evidence Alignment","Score","Original Verdict","Management Outcome","Priority","Required Next Action","Action Owner","Due Date","Status","Outcome Rationale"]
    at.append(headers)
    for i in items:
        at.append([_safe(i.get("id")),_safe(i.get("remark_date")),_safe(i.get("retailer_code")),_safe(i.get("retailer_name")),_safe(i.get("submitted_role")),_safe(i.get("submitted_by")),_safe(i.get("kam")),_safe(i.get("ss")),_safe(i.get("tl")),_safe(i.get("rds")),_safe(i.get("club")),_safe(i.get("town_name")),_safe(i.get("district_name")),i["target_volume"],_safe(i.get("remark")),_safe(i["issue"]),_safe(i["evidence"]),i["score"],_safe(i["verdict"]),_safe(i["decision"]),_safe(i["priority"]),_safe(i["next_action"]),_safe(i["owner"]),i["due"],"Open",_safe(i["justification"])])
    _style_header(at); _set_widths(at,[7,12,15,34,8,20,18,18,18,22,13,14,18,11,48,22,32,9,26,28,11,48,24,14,12,55]); at.freeze_panes="A2"; at.auto_filter.ref=at.dimensions; _add_table(at,"EODOutcomeActions")
    for row in at.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)

    # Detailed Review
    dr=wb.create_sheet("Detailed Review")
    dh=["ID","Date","Retailer Code","Retailer","Role","Submitted By","KAM","SS","TL","RDS","Club","Town","District","Target Volume","Last Visit","Visit Count","EOD Remark","Issue Category","Evidence Alignment","Score","Verdict","Action Plan?","Recent Visit?","Justification","Management Recommendation"]
    dr.append(dh)
    for i in items:
        dr.append([_safe(i.get("id")),_safe(i.get("remark_date")),_safe(i.get("retailer_code")),_safe(i.get("retailer_name")),_safe(i.get("submitted_role")),_safe(i.get("submitted_by")),_safe(i.get("kam")),_safe(i.get("ss")),_safe(i.get("tl")),_safe(i.get("rds")),_safe(i.get("club")),_safe(i.get("town_name")),_safe(i.get("district_name")),i["target_volume"],_safe(i.get("last_visit")),int(i.get("visit_count") or 0),_safe(i.get("remark")),_safe(i["issue"]),_safe(i["evidence"]),i["score"],_safe(i["verdict"]),"Yes" if i["action"] else "No","Yes" if i["recent"] else "No",_safe(i["justification"]),_safe(i["recommendation"])])
    _style_header(dr); _set_widths(dr,[7,12,15,34,8,20,18,18,18,22,13,14,18,12,13,10,48,22,32,9,28,12,12,55,55]); dr.freeze_panes="A2"; dr.auto_filter.ref=dr.dimensions; _add_table(dr,"EODDetailedReview")
    for row in dr.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)

    # Submitter Scorecard
    sc=wb.create_sheet("Submitter Scorecard")
    sh=["Role","Submitted By","Remarks","Strong","Partial","Weak","Avg Score","Action-ready","Recent Visit","High/Critical","Management Comment"]
    sc.append(sh)
    groups=defaultdict(list)
    for i in items: groups[(i.get("submitted_role") or "",i.get("submitted_by") or "")].append(i)
    for (role,name),grp in sorted(groups.items()):
        strong_n=sum(x["verdict"].startswith("Strongly justified") for x in grp); part_n=sum(x["verdict"]=="Partially justified" for x in grp); weak_n=sum(x["verdict"].startswith("Weak") for x in grp)
        avg=round(sum(x["score"] for x in grp)/len(grp),1); act_n=sum(x["action"] for x in grp); recent_n=sum(x["recent"] for x in grp); hi=sum(x["priority"] in ("Critical","High") for x in grp)
        comment="Good decision-grade remarks" if weak_n==0 and avg>=60 else "Mixed quality; strengthen evidence/action" if weak_n < len(grp)/2 else "High rework need; too many basic explanations"
        sc.append([role,name,len(grp),strong_n,part_n,weak_n,avg,act_n,recent_n,hi,comment])
    _style_header(sc); _set_widths(sc,[10,24,10,10,10,10,12,12,12,14,42]); sc.freeze_panes="A2"; _add_table(sc,"EODSubmitterScorecard")

    # Issue Analysis
    ia=wb.create_sheet("Issue Analysis")
    ia.append(["Issue Category","Remarks","%","Avg Score","Weak","High/Critical","Management Interpretation"])
    for issue,count in sorted(ic.items(),key=lambda x:(-x[1],x[0])):
        grp=[i for i in items if i["issue"]==issue]; avg=round(sum(x["score"] for x in grp)/len(grp),1)
        ia.append([issue,count,count/total if total else 0,avg,sum(x["verdict"].startswith("Weak") for x in grp),sum(x["priority"] in ("Critical","High") for x in grp),interp.get(issue,"Needs management validation.")])
    _style_header(ia); _set_widths(ia,[25,10,10,12,10,14,55]); ia.freeze_panes="A2"; _add_table(ia,"EODIssueAnalysis")
    for c in ia["C"][1:]: c.number_format="0.0%"

    # Existing coverage data retained
    if coverage is not None:
        cov=wb.create_sheet("Submission Coverage")
        cov.append(["Date","Role","Name","Eligible Retailers","Remarks Submitted","Remarks Pending"])
        for x in coverage:
            cov.append([x["date"],x["role"],x["name"],x["eligible_count"] if x["eligible_count"] is not None else "Unavailable",x["submitted_count"],x["pending_count"] if x["pending_count"] is not None else "Unavailable"])
        _style_header(cov); _set_widths(cov,[14,10,30,20,20,20]); cov.freeze_panes="A2"; _add_table(cov,"EODSubmissionCoverage")

    meth=wb.create_sheet("Methodology")
    meth.append(["Item","Definition / Limitation"])
    method_rows=[
      ("Purpose","Assess whether EOD remarks identify a plausible retailer-specific cause and convert them into an actionable management outcome."),
      ("Strongly justified","Specific root cause with evidence/action, or confirmed structural/master-data issue."),
      ("Partially justified","Plausible issue but evidence, quantification, visit support or closure plan is incomplete."),
      ("Weak / basic","Mostly repeats low sale/footfall or makes an unsupported promise without an evidenced root cause."),
      ("Action-ready","Remark contains a follow-up, discussion, replacement, promotion, billing, commitment or other concrete next step."),
      ("Recent visit","Latest recorded visit is same day or previous day relative to the remark date."),
      ("Important limitation","The EOD table does not preserve the exact historical stock/DOS and FTD/MTD snapshot at submission time. Scoring therefore evaluates remark quality, master-data context and visit evidence; it does not claim every sales/stock statement is independently verified."),
      ("Management rule","Do not accept 'sale low', 'footfall low' or 'sale will come' as complete reasons without root cause, evidence and next action.")
    ]
    for r in method_rows: meth.append(r)
    _style_header(meth); _set_widths(meth,[24,100])
    for row in meth.iter_rows(): 
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)

    # Basic charts on Management Outcome
    try:
        chart=BarChart(); data=Reference(ws,min_col=2,min_row=14,max_row=20); cats=Reference(ws,min_col=1,min_row=15,max_row=20)
        chart.add_data(data,titles_from_data=True); chart.set_categories(cats); chart.title="Management Outcome Distribution"; chart.height=7; chart.width=11; chart.legend=None; ws.add_chart(chart,"I14")
    except Exception:
        pass

    out=io.BytesIO(); wb.save(out); out.seek(0); return out
