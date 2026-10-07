import io
from datetime import date
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app.eod_export import build_eod_workbook
from app.main import app, crud
from app.auth import get_dashboard_user


def test_workbook_has_all_submitters_and_literal_remarks():
    rows = [
        {"submitted_role": "TL", "submitted_by": "TL One", "remark": "=1+1"},
        {"submitted_role": "SS", "submitted_by": "SS Two", "remark": "Follow up tomorrow"},
        {"submitted_role": "TL", "submitted_by": "TL Three", "remark": "Stock pending"},
    ]
    wb = load_workbook(build_eod_workbook(rows, date(2026, 10, 1), date(2026, 10, 7)))
    assert wb["All Remarks"].max_row == 4
    assert wb["TL Remarks"].max_row == 3
    assert wb["SS Remarks"].max_row == 2
    assert wb["All Remarks"]["I2"].value == "=1+1"
    assert wb["All Remarks"]["I2"].data_type == "s"


def test_empty_report_still_has_headers():
    wb = load_workbook(build_eod_workbook([], date(2026, 10, 7), date(2026, 10, 7)))
    assert wb["All Remarks"].max_row == 1
    assert wb["TL Remarks"]["B1"].value == "Submitted By"


def test_export_requires_management_and_valid_dates(monkeypatch):
    calls = []
    monkeypatch.setattr(crud, "list_management_eod_remarks", lambda start, end: calls.append((start, end)) or [])
    client = TestClient(app)
    try:
        assert client.get("/exports/eod-remarks.xlsx").status_code == 401
        app.dependency_overrides[get_dashboard_user] = lambda: {"role": "TL", "name": "TL One"}
        assert client.get("/exports/eod-remarks.xlsx").status_code == 403
        assert calls == []
        app.dependency_overrides[get_dashboard_user] = lambda: {"role": "MANAGER", "name": "Manager"}
        assert client.get("/exports/eod-remarks.xlsx?date_from=2026-10-07&date_to=2026-10-01").status_code == 400
        response = client.get("/exports/eod-remarks.xlsx?date_from=2026-10-01&date_to=2026-10-07")
        assert response.status_code == 200
        assert calls == [("2026-10-01", "2026-10-07")]
        assert "eod-remarks-tl-ss-2026-10-01-to-2026-10-07.xlsx" in response.headers["content-disposition"]
        assert load_workbook(io.BytesIO(response.content))["All Remarks"].max_row == 1
    finally:
        app.dependency_overrides.pop(get_dashboard_user, None)
