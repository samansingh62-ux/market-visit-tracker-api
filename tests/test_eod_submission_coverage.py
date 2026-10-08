import ast
import calendar
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from app.eod_export import build_eod_workbook


def coverage_context():
    tree = ast.parse(Path("app/main.py").read_text())
    funcs = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in
             ("_eod_submission_coverage", "_eod_is_low_sale")]
    accounts = [{"role": "TL", "name": "TL1"}, {"role": "SS", "name": "=SS2"}]
    crud = SimpleNamespace(
        list_active_users_by_role=lambda role: [u for u in accounts if u["role"] == role],
        get_performance_scope_rows=lambda month, **scope: [
            {"code": "R1", "target_volume": 250}, {"code": "R2", "target_volume": 240}],
    )
    ns = dict(date=date, timedelta=timedelta, calendar=calendar, crud=crud,
              _scope=lambda u: {"name": u["name"]},
              _eod_sales_records=lambda start, end: {"R1": 6, "R2": 0} if start == end else {"R1": 10, "R2": 0},
              _eod_sales_totals=lambda records, scope: records)
    exec(compile(ast.Module(body=funcs, type_ignores=[]), "coverage", "exec"), ns)
    return ns


def test_summary_coverage_and_formula_safety():
    ns = coverage_context()
    day = date(2026, 10, 8)
    rows = [dict(remark_date=day.isoformat(), submitted_by="TL1", submitted_role="TL",
                 retailer_code="R2", remark="=bad")]
    coverage = ns["_eod_submission_coverage"](rows, day, day)
    assert coverage[0]["eligible_count"] == 1
    assert coverage[0]["submitted_count"] == 1
    assert coverage[0]["pending_count"] == 0
    assert coverage[1]["eligible_count"] == 1
    assert coverage[1]["pending_count"] == 1
    workbook = load_workbook(build_eod_workbook(rows, day, day, coverage))
    values = list(workbook["Report Summary"].values)
    assert any(row[:2] == ("TLs who submitted remarks", 1) for row in values)
    assert any(row[:2] == ("SSs who submitted remarks", 0) for row in values)
    assert any(row[:4] == (day.isoformat(), "SS", "=SS2", 1) for row in values)
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    assert ns["_eod_is_low_sale"](6, 10) is False


def test_unavailable_sales_are_not_zero():
    ns = coverage_context()
    def unavailable(*args):
        raise RuntimeError("offline")
    ns["_eod_sales_records"] = unavailable
    day = date(2026, 10, 8)
    coverage = ns["_eod_submission_coverage"]([], day, day)
    assert all(row["eligible_count"] is None for row in coverage)
    workbook = load_workbook(build_eod_workbook([], day, day, coverage))
    assert any("Unavailable" in row for row in workbook["Report Summary"].values)
